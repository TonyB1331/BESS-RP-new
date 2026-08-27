"""Loader for the IPTO Imbalance Settlement Price Excel file.

Expected format (Sheet1):
    version | tradingPeriod           | 3-707-1
    1       | 2025-10-01 00:00:00     | 86.105
    1       | 2025-10-01 00:15:00     | 38.009
    ...

The column '3-707-1' contains the single imbalance settlement price
(€/MWh) per 15-minute ISP.  This is what the supplier actually
pays/receives at settlement — distinct from BM Up/Down prices.

Usage:
    loader = ImbalancePriceLoader("Imbalance_Price.xlsx")
    loader.load()
    prices = loader.get_day(date(2025, 10, 1))  # → np.array shape (96,)
"""

from __future__ import annotations

import logging
import re
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional

import numpy as np

logger = logging.getLogger(__name__)


class ImbalancePriceLoader:

    def __init__(self, path: str | Path) -> None:
        self._path = Path(path)
        self._data: Dict[date, np.ndarray] = {}

    def load(self) -> None:
        import openpyxl

        wb = openpyxl.load_workbook(str(self._path), data_only=True)
        ws = wb[wb.sheetnames[0]]

        raw: Dict[date, List[float]] = {}

        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) < 3:
                continue
            ts_raw, price = row[1], row[2]
            if ts_raw is None or price is None:
                continue

            # Handle DST "(2)" suffix
            ts_str = str(ts_raw).strip()
            ts_str = re.sub(r"\s*\(\d+\)\s*$", "", ts_str)
            try:
                dt = datetime.fromisoformat(ts_str)
            except ValueError:
                continue

            d = dt.date()
            if d not in raw:
                raw[d] = []
            raw[d].append(float(price))

        wb.close()

        # Convert to fixed-length arrays (96 slots)
        for d, prices in raw.items():
            self._data[d] = self._fit_length(np.array(prices, dtype=float), 96)

        logger.info(
            "Loaded imbalance prices for %d days (%s → %s)",
            len(self._data),
            min(self._data.keys()) if self._data else "?",
            max(self._data.keys()) if self._data else "?",
        )

    @staticmethod
    def _fit_length(arr: np.ndarray, target: int) -> np.ndarray:
        """Pad or truncate array to exactly `target` elements."""
        if len(arr) == target:
            return arr
        if len(arr) > target:
            return arr[:target]
        # Pad with last value
        return np.pad(arr, (0, target - len(arr)), mode="edge")

    def get_day(self, d: date) -> Optional[np.ndarray]:
        """Return 96-element array of imbalance prices for date d."""
        return self._data.get(d)

    def available_dates(self) -> List[date]:
        return sorted(self._data.keys())

    def __len__(self) -> int:
        return len(self._data)

    def __contains__(self, d: date) -> bool:
        return d in self._data
